import openpyxl

def getRowCount(file, sheetname):
    workbook = openpyxl.load_workbook(file)
    sheet = workbook[sheetname]
    row_count = sheet.max_row
    workbook.close()
    return row_count


def getColCount(file, sheetname):
    workbook = openpyxl.load_workbook(file)
    sheet = workbook[sheetname]
    col_count = sheet.max_column
    workbook.close()  # Close workbook after use
    return col_count

def readData(file, sheetname, rownum, colnum):
    workbook = openpyxl.load_workbook(file)
    sheet = workbook[sheetname]
    value = sheet.cell(row=rownum, column=colnum).value
    workbook.close()  # Close workbook after use
    return value

def writeData(file, sheetname, rownum, colnum, data):
    workbook = openpyxl.load_workbook(file)
    sheet= workbook[sheetname]
    sheet.cell(row=rownum, column=colnum, value=data)  # Corrected data assignment
    workbook.save(file)  # Save changes
    workbook.close()  # Properly close the workbook